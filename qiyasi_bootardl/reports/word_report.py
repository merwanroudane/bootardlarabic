"""Word (.docx) report generator with RTL Arabic content.

مولّد تقرير Word عربي. يتطلب python-docx.
"""
from __future__ import annotations

import pandas as pd

from ..utils.arabic_text import CASE_NAMES_AR, TERMS_AR, IC_NAMES_AR, REFERENCES
from ..utils.formatting import fmt, pct


def _set_rtl(paragraph) -> None:
    try:
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        pPr = paragraph._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
        paragraph.alignment = 2  # WD_ALIGN_PARAGRAPH.RIGHT
    except Exception:  # pragma: no cover
        pass


def _add_table(doc, df: pd.DataFrame, include_index: bool = True) -> None:
    ncols = len(df.columns) + (1 if include_index else 0)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    offset = 0
    if include_index:
        hdr[0].text = ""
        offset = 1
    for j, col in enumerate(df.columns):
        hdr[j + offset].text = str(col)
    for idx, row in df.iterrows():
        cells = table.add_row().cells
        if include_index:
            cells[0].text = str(idx)
        for j, col in enumerate(df.columns):
            cells[j + offset].text = fmt(row[col]) if isinstance(row[col], float) else str(row[col])


def build_word_report(result, path: str) -> str:
    """Write a .docx report and return ``path``."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "تصدير Word يتطلب python-docx. ثبّته عبر: pip install python-docx"
        ) from exc

    r = result
    doc = Document()

    h = doc.add_heading(TERMS_AR["test_name"], level=0)
    _set_rtl(h)

    for label, value in [
        (TERMS_AR["dependent"], r.yvar),
        (TERMS_AR["independent"], "، ".join(r.xvar)),
        (TERMS_AR["selected_case"], CASE_NAMES_AR.get(r.case, r.case)),
        (TERMS_AR["n_obs"], r.n_obs),
        (TERMS_AR["ardl_ic"], IC_NAMES_AR.get(r.ardl_ic, r.ardl_ic)),
        (TERMS_AR["selected_lags"], str(r.diff_lags)),
        (TERMS_AR["n_boot"], f"{r.n_boot} (فعّالة {r.bootstrap.n_boot_effective})"),
    ]:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value))
        _set_rtl(p)

    h = doc.add_heading(f"{TERMS_AR['final_decision']} (عند {pct(r.decision_level)})", level=1)
    _set_rtl(h)
    p = doc.add_paragraph()
    p.add_run(r.decision.label_ar).bold = True
    _set_rtl(p)
    p = doc.add_paragraph(r.decision.detail_ar)
    _set_rtl(p)

    h = doc.add_heading(TERMS_AR["interpretation"], level=1)
    _set_rtl(h)
    p = doc.add_paragraph(r.تفسير())
    _set_rtl(p)

    h = doc.add_heading(TERMS_AR["test_statistics"], level=1)
    _set_rtl(h)
    _add_table(doc, r.جدول_الإحصائيات())

    h = doc.add_heading(TERMS_AR["boot_crit"], level=1)
    _set_rtl(h)
    _add_table(doc, r.جدول_البوتستراب())

    h = doc.add_heading(TERMS_AR["boot_pval"], level=1)
    _set_rtl(h)
    _add_table(doc, r.جدول_القيم_الاحتمالية())

    if r.warnings_list:
        h = doc.add_heading(TERMS_AR["warnings"], level=1)
        _set_rtl(h)
        for w in r.warnings_list:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"[{w.severity}] {w.message_ar} ").bold = True
            p.add_run(f"← {w.advice_ar}")
            _set_rtl(p)

    h = doc.add_heading("المراجع", level=1)
    _set_rtl(h)
    for ref in REFERENCES:
        p = doc.add_paragraph(ref, style="List Number")
        _set_rtl(p)

    doc.save(path)
    return path
